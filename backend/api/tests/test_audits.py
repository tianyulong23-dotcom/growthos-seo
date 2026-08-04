import asyncio
import io
import json
from datetime import UTC, datetime
from types import SimpleNamespace
from typing import Any

from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook

from app.api.routes.audits import get_audit_service
from app.core.config import Settings
from app.main import app
from app.modules.audit.service import (
    AuditProjectRecord,
    AuditIssueGroupRecord,
    AuditService,
    WorkflowState,
    group_issue_records,
    normalized_run_status,
    progress_percentage,
    run_response,
    spreadsheet_cell,
    spreadsheet_row_variants,
)
from app.modules.crawling.models import AuditIssue, CrawlRun, PageSpeedResult


class FakeWorkflowController:
    def __init__(self, error: Exception | None = None) -> None:
        self.error = error
        self.task: dict[str, Any] | None = None
        self.recalculation_task: dict[str, Any] | None = None
        self.workflow_id = ""
        self.signals: list[tuple[str, str]] = []
        self.start_count = 0
        self.signal_observer: Any = None
        self.workflow_states: dict[str, WorkflowState] = {}

    async def start(self, task: dict[str, Any], workflow_id: str) -> None:
        self.start_count += 1
        self.task = task
        self.workflow_id = workflow_id
        if self.error is not None:
            raise self.error

    async def start_recalculation(
        self,
        task: dict[str, Any],
        workflow_id: str,
    ) -> None:
        self.start_count += 1
        self.recalculation_task = task
        self.workflow_id = workflow_id
        if self.error is not None:
            raise self.error

    async def signal(self, workflow_id: str, signal_name: str) -> None:
        self.signals.append((workflow_id, signal_name))
        if self.signal_observer is not None:
            self.signal_observer(signal_name)
        if self.error is not None:
            raise self.error

    async def signal_and_wait(
        self,
        workflow_id: str,
        signal_name: str,
        timeout_seconds: float,
    ) -> None:
        assert timeout_seconds > 0
        await self.signal(workflow_id, signal_name)

    async def status(self, workflow_id: str) -> WorkflowState:
        if self.error is not None:
            raise self.error
        return self.workflow_states.get(workflow_id, WorkflowState.RUNNING)


class FakeAuditObjectCleaner:
    def __init__(self, error: Exception | None = None) -> None:
        self.error = error
        self.calls: list[tuple[str, str, str]] = []

    async def delete_run_objects(
        self,
        organization_id: str,
        project_id: str,
        run_id: str,
    ) -> None:
        self.calls.append((organization_id, project_id, run_id))
        if self.error is not None:
            raise self.error


def test_partial_status_is_terminal_for_shared_status_progress() -> None:
    assert normalized_run_status("partial") == "partial"
    assert progress_percentage("completed", 5, 5, "partial") == 100


def test_post_crawl_stages_have_distinct_progress() -> None:
    assert progress_percentage("selecting_pages", 20, 20, "running", 20) == 91
    assert progress_percentage("checking_links", 20, 20, "running", 20) == 92
    assert progress_percentage("checking_links", 20, 20, "paused", 20) == 92
    assert progress_percentage("generating_profile", 5, 5, "running", 5) == 96
    assert progress_percentage("pagespeed", 20, 20, "running", 20) == 97
    assert progress_percentage("completed", 20, 20, "running", 20) == 99


def test_crawl_progress_uses_page_limit_without_regressing() -> None:
    early = progress_percentage("extracting_pages", 203, 51, "running", 1000)
    more_discovered = progress_percentage(
        "extracting_pages",
        459,
        51,
        "running",
        1000,
    )
    more_processed = progress_percentage(
        "extracting_pages",
        459,
        101,
        "running",
        1000,
    )

    assert more_discovered == early
    assert more_processed > more_discovered


def test_partial_technical_audit_response_is_exposed_as_completed() -> None:
    run = CrawlRun(
        run_id="partial-run",
        organization_id="test-org",
        project_id="example",
        task_type="technical_audit",
        status="partial",
        stage="completed",
        message="技术审计已部分完成",
        summary={
            "page_count": 5,
            "health_score": 80,
            "errors": 1,
            "warnings": 2,
            "notices": 3,
            "rendered_pages": 0,
        },
        created_at=datetime.now(UTC),
    )

    response = run_response(run)

    assert response.status == "completed"
    assert response.progress == 100
    assert response.summary is not None
    assert response.summary.page_count == 5


class FakeAuditRunRepository:
    def __init__(self) -> None:
        self.runs: dict[tuple[str, str, str], CrawlRun] = {}
        self.projects: dict[tuple[str, str], AuditProjectRecord] = {
            ("test-org", "example"): AuditProjectRecord(
                id="example",
                organization_id="test-org",
                domain="Example.com",
                country="US",
                language="en",
                audit_run_id=None,
            )
        }
        self.issues: dict[str, list[AuditIssue]] = {}
        self.pagespeed: dict[str, list[PageSpeedResult]] = {}
        self.health: dict[tuple[str, str], int] = {}
        self.issue_queries: list[tuple[int, int]] = []
        self.issue_group_queries: list[tuple[int, int]] = []
        self.page_queries: list[tuple[int, int, str, int | None, str | None]] = []
        self.link_queries: list[tuple[int, int, str, bool | None, str | None]] = []
        self.page_rows: dict[str, list[tuple[Any, int]]] = {}
        self.link_rows: dict[str, list[tuple[Any, str]]] = {}
        self.external_resource_rows: dict[str, list[Any]] = {}
        self.status_code_rows: dict[str, list[tuple[int, str, int]]] = {}
        self.visualization_rows: dict[
            str,
            tuple[list[Any], list[Any], int, int],
        ] = {}
        self.checkpoint_run_ids: set[str] = set()
        self.checkpoints: dict[str, dict[str, Any]] = {}
        self.dispatches: dict[str, tuple[str, dict[str, Any]]] = {}

    async def get_project(
        self,
        organization_id: str,
        project_id: str,
    ) -> AuditProjectRecord | None:
        return self.projects.get((organization_id, project_id))

    async def create_for_project(
        self,
        project: AuditProjectRecord,
        run: CrawlRun,
        task: dict[str, Any],
    ) -> bool:
        active_statuses = {"queued", "running", "stopping"}
        if any(
            item.organization_id == project.organization_id
            and item.project_id == project.id
            and item.task_type == "technical_audit"
            and item.archived_at is None
            and normalized_run_status(item.status) in active_statuses
            for item in self.runs.values()
        ):
            return False
        self.runs[(run.organization_id, run.project_id, run.run_id)] = run
        self.dispatches[run.run_id] = (str(run.temporal_workflow_id), task)
        self.projects[(project.organization_id, project.id)] = AuditProjectRecord(
            id=project.id,
            organization_id=project.organization_id,
            domain=project.domain,
            country=project.country,
            language=project.language,
            audit_run_id=run.run_id,
        )
        return True

    async def list_pending_dispatches(
        self,
        limit: int,
    ) -> list[tuple[str, str, dict[str, Any]]]:
        return [
            (run_id, workflow_id, task)
            for run_id, (workflow_id, task) in list(self.dispatches.items())[:limit]
        ]

    async def mark_dispatch_succeeded(self, run_id: str) -> None:
        return None

    async def record_dispatch_failure(self, run_id: str, message: str) -> None:
        run = next(
            run for (_, _, item_run_id), run in self.runs.items() if item_run_id == run_id
        )
        run.message = "任务已保存，等待技术审计服务恢复后自动重试"

    async def get_run(
        self,
        organization_id: str,
        project_id: str,
        run_id: str,
    ) -> CrawlRun | None:
        return self.runs.get((organization_id, project_id, run_id))

    async def list_runs(
        self,
        organization_id: str,
        project_id: str,
        include_archived: bool,
        page: int,
        page_size: int,
        search: str,
        status: str | None,
    ) -> tuple[list[CrawlRun], int]:
        runs = [
            run
            for (org_id, item_project_id, _), run in self.runs.items()
            if org_id == organization_id and item_project_id == project_id
        ]
        if not include_archived:
            runs = [run for run in runs if run.archived_at is None]
        if status:
            statuses = {"completed", "partial"} if status == "completed" else {status}
            runs = [run for run in runs if run.status in statuses]
        if search:
            query = search.casefold()
            runs = [
                run
                for run in runs
                if query
                in " ".join(
                    [run.run_id, run.stage or "", run.message or ""]
                ).casefold()
            ]
        runs = sorted(runs, key=lambda run: run.created_at, reverse=True)
        start = (page - 1) * page_size
        return runs[start : start + page_size], len(runs)

    async def list_active_runs(self) -> list[CrawlRun]:
        return [
            run
            for run in self.runs.values()
            if normalized_run_status(run.status) in {"queued", "running", "stopping"}
        ]

    async def update_run_if_status(
        self,
        run_id: str,
        expected_statuses: set[str],
        *,
        expected_workflow_id: str | None = None,
        **values: Any,
    ) -> bool:
        run = next(run for (_, _, item_run_id), run in self.runs.items() if item_run_id == run_id)
        if run.status not in expected_statuses:
            return False
        if expected_workflow_id is not None and run.temporal_workflow_id != expected_workflow_id:
            return False
        for key, value in values.items():
            setattr(run, key, value)
        return True

    async def resume_run(
        self,
        project: AuditProjectRecord,
        run_id: str,
        expected_statuses: set[str],
        **values: Any,
    ) -> bool:
        active_statuses = {"queued", "running", "stopping"}
        if any(
            item.organization_id == project.organization_id
            and item.project_id == project.id
            and item.run_id != run_id
            and item.task_type == "technical_audit"
            and item.archived_at is None
            and normalized_run_status(item.status) in active_statuses
            for item in self.runs.values()
        ):
            return False
        updated = await self.update_run_if_status(
            run_id,
            expected_statuses,
            **values,
        )
        if not updated:
            return False
        self.projects[(project.organization_id, project.id)] = AuditProjectRecord(
            id=project.id,
            organization_id=project.organization_id,
            domain=project.domain,
            country=project.country,
            language=project.language,
            audit_run_id=run_id,
        )
        return True

    async def archive_run(
        self,
        project: AuditProjectRecord,
        run_id: str,
        archived_at: datetime,
    ) -> bool:
        run = self.runs[(project.organization_id, project.id, run_id)]
        if normalized_run_status(run.status) in {"queued", "running", "stopping"}:
            return False
        run.archived_at = archived_at
        current = self.projects[(project.organization_id, project.id)]
        if current.audit_run_id == run_id:
            replacements = [
                item
                for item in self.runs.values()
                if item.organization_id == project.organization_id
                and item.project_id == project.id
                and item.run_id != run_id
                and item.task_type == "technical_audit"
                and item.archived_at is None
            ]
            replacement = max(replacements, key=lambda item: item.created_at, default=None)
            self.projects[(project.organization_id, project.id)] = AuditProjectRecord(
                id=project.id,
                organization_id=project.organization_id,
                domain=project.domain,
                country=project.country,
                language=project.language,
                audit_run_id=replacement.run_id if replacement is not None else None,
            )
        return True

    async def has_checkpoint(self, run_id: str) -> bool:
        return run_id in self.checkpoint_run_ids

    async def load_checkpoint(self, run_id: str) -> dict[str, Any] | None:
        return self.checkpoints.get(run_id)

    async def update_project_health(
        self,
        organization_id: str,
        project_id: str,
        run_id: str,
        health: int,
    ) -> None:
        project = self.projects[(organization_id, project_id)]
        if project.audit_run_id == run_id:
            self.health[(organization_id, project_id)] = health

    async def delete_run(
        self,
        project: AuditProjectRecord,
        run_id: str,
    ) -> bool:
        run = self.runs.get((project.organization_id, project.id, run_id))
        if run is None:
            return False
        if normalized_run_status(run.status) in {"queued", "running", "stopping"}:
            return False
        self.runs.pop((project.organization_id, project.id, run_id), None)
        current = self.projects[(project.organization_id, project.id)]
        if current.audit_run_id == run_id:
            replacements = [
                item
                for item in self.runs.values()
                if item.organization_id == project.organization_id
                and item.project_id == project.id
                and item.task_type == "technical_audit"
                and item.archived_at is None
            ]
            replacement = max(replacements, key=lambda item: item.created_at, default=None)
            self.projects[(project.organization_id, project.id)] = AuditProjectRecord(
                id=project.id,
                organization_id=project.organization_id,
                domain=project.domain,
                country=project.country,
                language=project.language,
                audit_run_id=replacement.run_id if replacement is not None else None,
            )
        return True

    async def list_pagespeed(self, run_id: str) -> list[PageSpeedResult]:
        return self.pagespeed.get(run_id, [])

    async def pagespeed_counts(self, run_ids: list[str]) -> dict[str, tuple[int, int]]:
        counts: dict[str, tuple[int, int]] = {}
        for run_id in run_ids:
            results = self.pagespeed.get(run_id, [])
            counts[run_id] = (len(results), sum(not result.error for result in results))
        return counts

    async def list_issues(self, run_id: str) -> list[AuditIssue]:
        return self.issues.get(run_id, [])

    async def list_issue_groups(
        self,
        run_id: str,
        page: int,
        page_size: int,
        severity: str | None,
        search: str,
    ) -> tuple[list[AuditIssueGroupRecord], int, bool]:
        issues = self.issues.get(run_id, [])
        if not issues:
            return [], 0, False
        self.issue_group_queries.append((page, page_size))
        groups = group_issue_records(issues, severity, search)
        start = (page - 1) * page_size
        return groups[start : start + page_size], len(groups), True

    async def list_issues_batch(
        self,
        run_id: str,
        page: int,
        page_size: int,
    ) -> list[AuditIssue]:
        self.issue_queries.append((page, page_size))
        rows = self.issues.get(run_id, [])
        start = (page - 1) * page_size
        return rows[start : start + page_size]

    async def list_pages(
        self,
        run_id: str,
        page: int,
        page_size: int,
        search: str,
        status_code: int | None,
        status_family: str | None,
    ) -> tuple[list[Any], int]:
        self.page_queries.append((page, page_size, search, status_code, status_family))
        rows = self.page_rows.get(run_id, [])
        start = (page - 1) * page_size
        return rows[start : start + page_size], len(rows)

    async def list_links(
        self,
        run_id: str,
        page: int,
        page_size: int,
        search: str,
        internal: bool | None,
        status_family: str | None,
    ) -> tuple[list[Any], int]:
        self.link_queries.append((page, page_size, search, internal, status_family))
        rows = self.link_rows.get(run_id, [])
        start = (page - 1) * page_size
        return rows[start : start + page_size], len(rows)

    async def list_external_resources(
        self,
        run_id: str,
        page: int,
        page_size: int,
        search: str,
        status_family: str | None,
    ) -> tuple[list[Any], int]:
        rows = self.external_resource_rows.get(run_id, [])
        start = (page - 1) * page_size
        return rows[start : start + page_size], len(rows)

    async def status_codes(self, run_id: str) -> list[tuple[int, str, int]]:
        return self.status_code_rows.get(run_id, [])

    async def visualization(
        self,
        run_id: str,
    ) -> tuple[list[Any], list[Any], int, int]:
        return self.visualization_rows.get(run_id, ([], [], 0, 0))

    def complete(self, run_id: str, summary: dict[str, int]) -> None:
        run = next(run for (_, _, item_run_id), run in self.runs.items() if item_run_id == run_id)
        run.status = "completed"
        run.stage = "completed"
        run.message = "技术审计已完成"
        run.discovered = summary["page_count"]
        run.processed = summary["page_count"]
        run.selected = summary["page_count"]
        run.page_count = summary["page_count"]
        run.summary = summary
        run.can_resume = False
        run.finished_at = datetime.now(UTC)


def build_service(
    *,
    launch_error: Exception | None = None,
    control_timeout_seconds: float = 30,
    object_cleaner: FakeAuditObjectCleaner | None = None,
) -> tuple[AuditService, FakeWorkflowController, FakeAuditRunRepository]:
    controller = FakeWorkflowController(launch_error)
    repository = FakeAuditRunRepository()
    return (
        AuditService(
            settings=Settings(
                app_env="test",
                default_organization_id="test-org",
                audit_control_timeout_seconds=control_timeout_seconds,
                audit_checkpoint_poll_interval_seconds=0.001,
            ),
            controller=controller,
            repository=repository,
            object_cleaner=object_cleaner,
        ),
        controller,
        repository,
    )


async def api_request(
    method: str,
    path: str,
    *,
    json: dict[str, Any] | None = None,
) -> tuple[int, Any]:
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        response = await client.request(method, path, json=json)
        body = response.json() if response.content else None
        return response.status_code, body


async def api_raw_request(
    method: str,
    path: str,
) -> tuple[int, str, str]:
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        response = await client.request(method, path)
        return (
            response.status_code,
            response.headers.get("content-type", ""),
            response.text,
        )


async def api_binary_request(
    method: str,
    path: str,
) -> tuple[int, str, str, bytes]:
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        response = await client.request(method, path)
        return (
            response.status_code,
            response.headers.get("content-type", ""),
            response.headers.get("content-disposition", ""),
            response.content,
        )


def test_create_audit_run_starts_crawler_workflow() -> None:
    service, controller, _ = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        status_code, body = asyncio.run(
            api_request(
                "POST",
                "/api/v1/projects/example/audit-runs",
                json={
                    "max_pages": 250,
                    "scope": "directory",
                    "directory": "/blog/",
                    "rendering": "off",
                    "allowed_paths": ["/blog/", "/guides/"],
                    "excluded_paths": ["/blog/private/"],
                    "ignored_parameters": ["utm_*", "gclid"],
                    "issue_exclusion_patterns": ["/admin/*", "*.json"],
                    "enable_duplication_check": False,
                    "duplication_threshold": 0.9,
                    "enable_pagespeed": True,
                },
            )
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 202
    assert body["status"] == "queued"
    assert controller.task is not None
    assert controller.task["target_url"] == "https://example.com"
    assert controller.task["directory"] == "/blog"
    assert controller.task["rendering"] == "off"
    assert controller.task["issue_exclusion_patterns"] == ["/admin/*", "*.json"]
    assert controller.task["enable_duplication_check"] is False
    assert controller.task["duplication_threshold"] == 0.9
    assert controller.task["enable_pagespeed"] is True
    assert controller.workflow_id == f"crawler:technical_audit:{body['run_id']}"


def test_create_rejects_second_active_audit_run() -> None:
    service, controller, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        first_status, first = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
        second_status, second = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
    finally:
        app.dependency_overrides.clear()

    assert first_status == 202
    assert second_status == 409
    assert second["detail"] == "这个项目已有正在运行的审计"
    assert len(repository.runs) == 1
    assert controller.start_count == 1
    assert next(iter(repository.runs.values())).run_id == first["run_id"]


def test_completed_audit_can_recalculate_issues_without_restarting_crawl() -> None:
    service, controller, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request(
                "POST",
                "/api/v1/projects/example/audit-runs",
                json={"issue_exclusion_patterns": ["/old/*"]},
            )
        )
        run_id = created["run_id"]
        repository.complete(
            run_id,
            {
                "page_count": 5,
                "health_score": 80,
                "errors": 1,
                "warnings": 2,
                "notices": 3,
                "rendered_pages": 0,
            },
        )

        status_code, body = asyncio.run(
            api_request(
                "POST",
                f"/api/v1/projects/example/audit-runs/{run_id}/recalculate-issues",
                json={"issue_exclusion_patterns": ["/new/*", "*.json"]},
            )
        )
    finally:
        app.dependency_overrides.clear()

    run = repository.runs[("test-org", "example", run_id)]
    assert status_code == 202
    assert body["status"] == "recalculating"
    assert body["progress"] == 99
    assert body["summary"]["page_count"] == 5
    assert run.config_snapshot["issue_exclusion_patterns"] == ["/old/*"]
    assert controller.recalculation_task is not None
    assert controller.recalculation_task["issue_exclusion_patterns"] == [
        "/new/*",
        "*.json",
    ]
    assert ":recalculate:" in controller.workflow_id


def test_recalculation_launch_failure_restores_completed_audit() -> None:
    service, _, repository = build_service(launch_error=RuntimeError("offline"))
    run = CrawlRun(
        run_id="00000000-0000-0000-0000-000000000001",
        organization_id="test-org",
        project_id="example",
        task_type="technical_audit",
        target_url="https://example.com",
        country="US",
        language="en",
        status="completed",
        stage="completed",
        message="Audit completed",
        config_snapshot={"issue_exclusion_patterns": ["/old/*"]},
        summary={
            "page_count": 1,
            "health_score": 100,
            "errors": 0,
            "warnings": 0,
            "notices": 0,
            "rendered_pages": 0,
        },
        created_at=datetime.now(UTC),
    )
    repository.runs[("test-org", "example", run.run_id)] = run
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        status_code, _ = asyncio.run(
            api_request(
                "POST",
                f"/api/v1/projects/example/audit-runs/{run.run_id}/recalculate-issues",
                json={"issue_exclusion_patterns": ["/new/*"]},
            )
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 503
    assert run.status == "completed"
    assert run.config_snapshot["issue_exclusion_patterns"] == ["/old/*"]


def test_get_completed_audit_run_returns_persisted_summary() -> None:
    service, _, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
        run_id = created["run_id"]
        repository.complete(
            run_id,
            {
                "page_count": 12,
                "health_score": 81,
                "errors": 3,
                "warnings": 5,
                "notices": 2,
                "rendered_pages": 4,
            },
        )
        status_code, body = asyncio.run(
            api_request("GET", f"/api/v1/projects/example/audit-runs/{run_id}")
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 200
    assert body["status"] == "completed"
    assert body["summary"]["page_count"] == 12
    assert body["summary"]["health_score"] == 81
    assert body["summary"]["errors"] == 3
    assert repository.health[("test-org", "example")] == 81


def test_activity_returns_incremental_checkpoint_pages() -> None:
    service, _, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
        run_id = created["run_id"]
        repository.checkpoints[run_id] = {
            "pages": [
                {
                    "url": "https://example.com/",
                    "final_url": "https://example.com/",
                    "status_code": 200,
                    "title": "Home",
                    "response_time_ms": 120,
                    "fetched_at": "2026-07-22T08:00:00Z",
                },
                {
                    "url": "https://example.com/missing",
                    "final_url": "https://example.com/missing",
                    "status_code": 404,
                    "title": "Missing",
                    "response_time_ms": 321,
                    "fetched_at": "2026-07-22T08:00:01Z",
                },
            ]
        }
        status_code, body = asyncio.run(
            api_request(
                "GET",
                f"/api/v1/projects/example/audit-runs/{run_id}/activity?cursor=1&limit=1",
            )
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 200
    assert body["next_cursor"] == 2
    assert body["items"] == [
        {
            "sequence": 2,
            "url": "https://example.com/missing",
            "final_url": "https://example.com/missing",
            "status_code": 404,
            "title": "Missing",
            "error": "",
            "error_type": "",
            "depth": None,
            "rendered": False,
            "response_time_ms": 321,
            "fetched_at": "2026-07-22T08:00:01Z",
        }
    ]


def test_reconcile_marks_closed_workflow_failed_and_resumable() -> None:
    service, controller, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
        run_id = created["run_id"]
        repository.checkpoint_run_ids.add(run_id)
        controller.workflow_states[controller.workflow_id] = WorkflowState.CLOSED

        reconciled = asyncio.run(service.reconcile_active_runs())
    finally:
        app.dependency_overrides.clear()

    run = repository.runs[("test-org", "example", run_id)]
    assert reconciled == 1
    assert run.status == "failed"
    assert run.stage == "failed"
    assert run.can_resume is True
    assert run.finished_at is not None


def test_reconcile_keeps_running_workflow_active() -> None:
    service, _, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
        run_id = created["run_id"]

        reconciled = asyncio.run(service.reconcile_active_runs())
    finally:
        app.dependency_overrides.clear()

    run = repository.runs[("test-org", "example", run_id)]
    assert reconciled == 0
    assert run.status == "queued"
    assert run.finished_at is None


def test_get_run_reconciles_workflow_that_closed_after_api_startup() -> None:
    service, controller, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
        run_id = created["run_id"]
        repository.checkpoint_run_ids.add(run_id)
        controller.workflow_states[controller.workflow_id] = WorkflowState.CLOSED

        status_code, body = asyncio.run(
            api_request("GET", f"/api/v1/projects/example/audit-runs/{run_id}")
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 200
    assert body["status"] == "failed"
    assert body["stage"] == "failed"
    assert body["can_resume"] is True


def test_issues_group_urls_from_persisted_issue_rows() -> None:
    service, _, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
        run_id = created["run_id"]
        repository.issues[run_id] = [
            AuditIssue(
                id=1,
                run_id=run_id,
                page_id=1,
                url="https://example.com/a",
                severity="warning",
                category="content",
                code="duplicate_content",
                issue="重复内容",
                details="页面内容与其他页面高度相似",
                related_url="https://example.com/b",
                similarity=0.91,
            ),
            AuditIssue(
                id=2,
                run_id=run_id,
                page_id=2,
                url="https://example.com/b",
                severity="warning",
                category="content",
                code="duplicate_content",
                issue="重复内容",
                details="页面内容与其他页面高度相似",
                related_url="https://example.com/a",
                similarity=0.91,
            ),
        ]
        status_code, body = asyncio.run(
            api_request(
                "GET",
                f"/api/v1/projects/example/audit-runs/{run_id}/issues",
            )
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 200
    assert body["total"] == 1
    assert body["items"][0]["affected_count"] == 2
    assert body["items"][0]["raw"]["max_similarity"] == 0.91


def test_issue_group_ids_are_unique_for_distinct_groups() -> None:
    service, _, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
        run_id = created["run_id"]
        repository.issues[run_id] = [
            AuditIssue(
                id=1,
                run_id=run_id,
                page_id=1,
                url="https://example.com/a",
                severity="error",
                category="performance",
                code="slow_response",
                issue="响应时间超过 3 秒",
                details="页面响应时间为 3100ms",
            ),
            AuditIssue(
                id=2,
                run_id=run_id,
                page_id=2,
                url="https://example.com/b",
                severity="error",
                category="performance",
                code="slow_response",
                issue="响应时间超过 5 秒",
                details="页面响应时间为 5200ms",
            ),
        ]
        status_code, body = asyncio.run(
            api_request(
                "GET",
                f"/api/v1/projects/example/audit-runs/{run_id}/issues",
            )
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 200
    assert body["total"] == 2
    assert len({item["id"] for item in body["items"]}) == 2


def test_pages_accepts_status_family_filter() -> None:
    service, _, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
        run_id = created["run_id"]
        status_code, body = asyncio.run(
            api_request(
                "GET",
                (
                    f"/api/v1/projects/example/audit-runs/{run_id}/pages"
                    "?page=2&page_size=25&search=missing&status_family=4xx"
                ),
            )
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 200
    assert body == {"items": [], "total": 0, "page": 2, "page_size": 25}
    assert repository.page_queries == [(2, 25, "missing", None, "4xx")]


def test_pages_rejects_invalid_status_family() -> None:
    service, _, _ = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
        run_id = created["run_id"]
        status_code, _ = asyncio.run(
            api_request(
                "GET",
                (f"/api/v1/projects/example/audit-runs/{run_id}/pages?status_family=6xx"),
            )
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 422


def test_pages_return_librecrawl_detail_fields() -> None:
    service, _, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
        run_id = created["run_id"]
        repository.page_rows[run_id] = [
            (
                SimpleNamespace(
                    page_id=1,
                    requested_url="https://example.com/",
                    final_url="https://example.com/",
                    status_code=200,
                    title="Example",
                    description="Example description",
                    content_type="text/html",
                    word_count=321,
                    response_time_ms=145,
                    rendered=True,
                    depth=0,
                    canonical="https://example.com/",
                    h1=["Example"],
                    h2=["Features"],
                    h3=[],
                    headings=["Example", "Features"],
                    meta_tags={"description": "Example description"},
                    size_bytes=4096,
                    language="en",
                    charset="utf-8",
                    viewport="width=device-width",
                    robots="index,follow",
                    author="Example Author",
                    keywords="example",
                    generator="Example CMS",
                    theme_color="#ffffff",
                    open_graph={"title": "Example"},
                    twitter_tags={"card": "summary"},
                    structured_data=[{"@type": "Organization"}],
                    schema_org=[{"type": "Organization"}],
                    analytics={"gtag": True},
                    images=[{"src": "/hero.jpg"}],
                    broken_images=[],
                    internal_links=8,
                    external_links=2,
                    hreflang=[{"hreflang": "en", "href": "https://example.com/"}],
                    redirects=[],
                    linked_from=["https://example.com/about"],
                    discovered_from="https://example.com/sitemap.xml",
                    error=None,
                    error_type=None,
                ),
                2,
            )
        ]
        status_code, body = asyncio.run(
            api_request(
                "GET",
                f"/api/v1/projects/example/audit-runs/{run_id}/pages",
            )
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 200
    page = body["items"][0]
    assert page["h1"] == ["Example"]
    assert page["size_bytes"] == 4096
    assert page["open_graph"] == {"title": "Example"}
    assert page["structured_data"] == [{"@type": "Organization"}]
    assert page["internal_links"] == 8
    assert page["external_links"] == 2
    assert page["issues_count"] == 2


def test_links_accept_status_filter_and_return_librecrawl_fields() -> None:
    service, _, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
        run_id = created["run_id"]
        repository.link_rows[run_id] = [
            (
                SimpleNamespace(
                    id=1,
                    target_url="https://external.example/resource",
                    anchor_text="Source",
                    target_status=404,
                    placement="footer",
                    is_internal=False,
                    rel="nofollow noopener",
                    target_domain="external.example",
                    in_navigation=True,
                ),
                "https://example.com/",
            )
        ]
        status_code, body = asyncio.run(
            api_request(
                "GET",
                (
                    f"/api/v1/projects/example/audit-runs/{run_id}/links"
                    "?internal=false&status_family=4xx"
                ),
            )
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 200
    assert repository.link_queries == [(1, 100, "", False, "4xx")]
    link = body["items"][0]
    assert link["placement"] == "footer"
    assert link["target_domain"] == "external.example"
    assert link["rel"] == "nofollow noopener"
    assert link["in_navigation"] is True
    assert link["follow"] is False


def test_status_codes_group_no_response_by_error_type() -> None:
    service, _, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
        run_id = created["run_id"]
        repository.status_code_rows[run_id] = [
            (200, "", 8),
            (404, "", 1),
            (0, "timeout", 1),
            (0, "dns_not_found", 2),
        ]
        status_code, body = asyncio.run(
            api_request(
                "GET",
                f"/api/v1/projects/example/audit-runs/{run_id}/status-codes",
            )
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 200
    assert body["total"] == 12
    assert body["items"][0]["percentage"] == 66.7
    statuses_by_error = {
        item["error_type"]: item["status"] for item in body["items"] if item["status_code"] == 0
    }
    assert statuses_by_error == {
        "timeout": "请求超时",
        "dns_not_found": "DNS 无法解析",
    }


def test_pause_sends_temporal_signal_and_updates_run() -> None:
    service, controller, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
        run_id = created["run_id"]
        run = repository.runs[("test-org", "example", run_id)]
        repository.checkpoint_run_ids.add(run_id)
        run.status = "running"
        run.discovered = 807
        run.processed = 1
        run.config_snapshot["max_pages"] = 30
        status_code, body = asyncio.run(
            api_request(
                "POST",
                f"/api/v1/projects/example/audit-runs/{run_id}/pause",
            )
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 200
    assert body["status"] == "paused"
    assert body["can_resume"] is True
    assert body["progress"] == 22
    assert controller.signals == [(run.temporal_workflow_id, "pause")]


def test_stop_sends_temporal_signal_and_marks_run_resumable() -> None:
    service, controller, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
        run_id = created["run_id"]
        run = repository.runs[("test-org", "example", run_id)]
        repository.checkpoint_run_ids.add(run_id)
        run.status = "running"
        run.discovered = 807
        run.processed = 1
        run.config_snapshot["max_pages"] = 30
        status_code, body = asyncio.run(
            api_request(
                "POST",
                f"/api/v1/projects/example/audit-runs/{run_id}/stop",
            )
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 200
    assert body["status"] == "stopped"
    assert body["can_resume"] is True
    assert body["progress"] == 22
    assert controller.signals == [(run.temporal_workflow_id, "stop")]


def test_pause_without_checkpoint_keeps_run_active() -> None:
    service, controller, repository = build_service(control_timeout_seconds=0.001)
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
        run_id = created["run_id"]
        run = repository.runs[("test-org", "example", run_id)]
        run.status = "running"
        status_code, _ = asyncio.run(
            api_request(
                "POST",
                f"/api/v1/projects/example/audit-runs/{run_id}/pause",
            )
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 409
    assert run.status == "running"
    assert run.can_resume is False
    assert controller.signals == []


def test_pause_does_not_overwrite_a_newer_terminal_state() -> None:
    service, controller, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
        run_id = created["run_id"]
        run = repository.runs[("test-org", "example", run_id)]
        run.status = "running"
        repository.checkpoint_run_ids.add(run_id)

        def finish_before_pause_update(_: str) -> None:
            run.status = "stopped"
            run.stage = "stopped"
            run.message = "审计已停止"

        controller.signal_observer = finish_before_pause_update
        status_code, body = asyncio.run(
            api_request(
                "POST",
                f"/api/v1/projects/example/audit-runs/{run_id}/pause",
            )
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 409
    assert body["detail"] == "审计状态已变化，请刷新后重试"
    assert run.status == "stopped"


def test_stop_persists_stopping_before_signaling_workflow() -> None:
    service, controller, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    observed_statuses: list[str] = []
    try:
        _, created = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
        run_id = created["run_id"]
        run = repository.runs[("test-org", "example", run_id)]
        run.status = "running"
        repository.checkpoint_run_ids.add(run_id)
        controller.signal_observer = lambda _: observed_statuses.append(run.status)

        status_code, body = asyncio.run(
            api_request(
                "POST",
                f"/api/v1/projects/example/audit-runs/{run_id}/stop",
            )
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 200
    assert observed_statuses == ["stopping"]
    assert body["status"] == "stopped"


def test_completed_stage_reports_99_until_result_is_persisted() -> None:
    service, _, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request(
                "POST",
                "/api/v1/projects/example/audit-runs",
                json={"max_pages": 30},
            )
        )
        run_id = created["run_id"]
        run = repository.runs[("test-org", "example", run_id)]
        run.status = "running"
        run.stage = "completed"
        run.discovered = 807
        run.processed = 30
        status_code, body = asyncio.run(
            api_request("GET", f"/api/v1/projects/example/audit-runs/{run_id}")
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 200
    assert body["status"] == "running"
    assert body["progress"] == 99


def test_resume_starts_a_new_workflow_from_saved_configuration() -> None:
    service, controller, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request(
                "POST",
                "/api/v1/projects/example/audit-runs",
                json={"max_pages": 321},
            )
        )
        run_id = created["run_id"]
        run = repository.runs[("test-org", "example", run_id)]
        run.status = "paused"
        run.can_resume = True
        repository.checkpoint_run_ids.add(run_id)
        original_workflow_id = controller.workflow_id
        status_code, body = asyncio.run(
            api_request(
                "POST",
                f"/api/v1/projects/example/audit-runs/{run_id}/resume",
            )
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 200
    assert body["status"] == "queued"
    assert body["can_resume"] is False
    assert controller.task is not None
    assert controller.task["run_id"] == run_id
    assert controller.task["max_pages"] == 321
    assert controller.workflow_id != original_workflow_id


def test_resume_missing_project_returns_not_found() -> None:
    service, _, repository = build_service()
    repository.projects.clear()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        status_code, _ = asyncio.run(
            api_request(
                "POST",
                ("/api/v1/projects/missing/audit-runs/00000000-0000-0000-0000-000000000000/resume"),
            )
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 404


def test_resume_rejects_stale_resumable_flag_without_checkpoint() -> None:
    service, controller, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
        run_id = created["run_id"]
        run = repository.runs[("test-org", "example", run_id)]
        run.status = "paused"
        run.stage = "paused"
        run.can_resume = True
        start_count = controller.start_count

        status_code, _ = asyncio.run(
            api_request(
                "POST",
                f"/api/v1/projects/example/audit-runs/{run_id}/resume",
            )
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 409
    assert controller.start_count == start_count
    assert run.status == "paused"


def test_resume_makes_the_resumed_run_current_for_the_project() -> None:
    service, _, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, first = asyncio.run(api_request("POST", "/api/v1/projects/example/audit-runs", json={}))
        first_run = repository.runs[("test-org", "example", first["run_id"])]
        first_run.status = "paused"
        first_run.stage = "paused"
        first_run.can_resume = True
        repository.checkpoint_run_ids.add(first["run_id"])

        _, second = asyncio.run(api_request("POST", "/api/v1/projects/example/audit-runs", json={}))
        repository.complete(
            second["run_id"],
            {
                "page_count": 1,
                "health_score": 90,
                "errors": 0,
                "warnings": 1,
                "notices": 0,
                "rendered_pages": 0,
            },
        )

        status_code, body = asyncio.run(
            api_request(
                "POST",
                f"/api/v1/projects/example/audit-runs/{first['run_id']}/resume",
            )
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 200
    assert body["status"] == "queued"
    assert repository.projects[("test-org", "example")].audit_run_id == first["run_id"]


def test_archive_is_hidden_from_default_history_and_delete_removes_run() -> None:
    cleaner = FakeAuditObjectCleaner()
    service, _, repository = build_service(object_cleaner=cleaner)
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
        run_id = created["run_id"]
        repository.complete(
            run_id,
            {
                "page_count": 1,
                "health_score": 100,
                "errors": 0,
                "warnings": 0,
                "notices": 0,
                "rendered_pages": 0,
            },
        )
        status_code, archived = asyncio.run(
            api_request(
                "POST",
                f"/api/v1/projects/example/audit-runs/{run_id}/archive",
            )
        )
        _, visible = asyncio.run(api_request("GET", "/api/v1/projects/example/audit-runs"))
        _, all_runs = asyncio.run(
            api_request(
                "GET",
                "/api/v1/projects/example/audit-runs?include_archived=true",
            )
        )
        delete_status, _ = asyncio.run(
            api_request(
                "DELETE",
                f"/api/v1/projects/example/audit-runs/{run_id}",
            )
        )
        get_status, _ = asyncio.run(
            api_request("GET", f"/api/v1/projects/example/audit-runs/{run_id}")
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 200
    assert archived["archived_at"] is not None
    assert visible["total"] == 0
    assert all_runs["total"] == 1
    assert delete_status == 204
    assert get_status == 404
    assert cleaner.calls == [("test-org", "example", run_id)]


def test_audit_history_supports_pagination_search_and_status_filtering() -> None:
    service, _, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        created_runs = []
        for index, status in enumerate(("completed", "failed", "partial"), start=1):
            _, created = asyncio.run(
                api_request("POST", "/api/v1/projects/example/audit-runs", json={})
            )
            run = repository.runs[("test-org", "example", created["run_id"])]
            run.status = status
            run.stage = "completed" if status in {"completed", "partial"} else "failed"
            run.message = f"history marker {index}"
            created_runs.append(run)

        page_status, page = asyncio.run(
            api_request(
                "GET",
                "/api/v1/projects/example/audit-runs?page=1&page_size=1",
            )
        )
        filter_status, completed = asyncio.run(
            api_request(
                "GET",
                "/api/v1/projects/example/audit-runs?status=completed",
            )
        )
        search_status, searched = asyncio.run(
            api_request(
                "GET",
                "/api/v1/projects/example/audit-runs?search=marker%202",
            )
        )
    finally:
        app.dependency_overrides.clear()

    assert page_status == 200
    assert page["page"] == 1
    assert page["page_size"] == 1
    assert page["total"] == 3
    assert len(page["items"]) == 1
    assert filter_status == 200
    assert completed["total"] == 2
    assert {item["status"] for item in completed["items"]} == {"completed"}
    assert search_status == 200
    assert searched["total"] == 1
    assert searched["items"][0]["run_id"] == created_runs[1].run_id


def test_delete_keeps_database_run_when_object_cleanup_fails() -> None:
    cleaner = FakeAuditObjectCleaner(RuntimeError("s3 unavailable"))
    service, _, repository = build_service(object_cleaner=cleaner)
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
        run_id = created["run_id"]
        repository.complete(
            run_id,
            {
                "page_count": 1,
                "health_score": 100,
                "errors": 0,
                "warnings": 0,
                "notices": 0,
                "rendered_pages": 0,
            },
        )

        status_code, _ = asyncio.run(
            api_request(
                "DELETE",
                f"/api/v1/projects/example/audit-runs/{run_id}",
            )
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 503
    assert cleaner.calls == [("test-org", "example", run_id)]
    assert ("test-org", "example", run_id) in repository.runs


def test_active_result_endpoints_fall_back_to_checkpoint() -> None:
    service, _, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
        run_id = created["run_id"]
        repository.checkpoints[run_id] = {
            "pages": [
                {
                    "url": "https://example.com/",
                    "final_url": "https://example.com/",
                    "status_code": 200,
                    "title": "Home",
                    "description": "Description",
                    "content_type": "text/html",
                    "h1": ["Home"],
                    "depth": 0,
                    "links": [
                        {
                            "url": "https://external.example/missing",
                            "text": "External",
                            "is_internal": False,
                            "target_status": 404,
                            "target_domain": "external.example",
                            "placement": "body",
                        }
                    ],
                }
            ],
            "issues": [
                {
                    "type": "error",
                    "category": "content",
                    "code": "missing_title",
                    "issue": "缺少标题",
                    "details": "页面没有 title",
                    "url": "https://example.com/",
                }
            ],
            "external_resources": [
                {
                    "url": "https://external.example/missing",
                    "final_url": "https://external.example/missing",
                    "status_code": 404,
                    "content_type": "text/html",
                    "size_bytes": 123,
                    "title": "Missing",
                    "checked_at": "2026-07-22T00:00:00Z",
                }
            ],
        }

        responses = {
            name: asyncio.run(
                api_request(
                    "GET",
                    f"/api/v1/projects/example/audit-runs/{run_id}/{path}",
                )
            )
            for name, path in {
                "pages": "pages",
                "issues": "issues",
                "links": "links?internal=false",
                "resources": "resources",
                "status_codes": "status-codes",
                "visualization": "visualization",
            }.items()
        }
    finally:
        app.dependency_overrides.clear()

    assert all(status_code == 200 for status_code, _ in responses.values())
    assert responses["pages"][1]["items"][0]["title"] == "Home"
    assert responses["pages"][1]["items"][0]["issues_count"] == 1
    assert responses["issues"][1]["items"][0]["code"] == "missing_title"
    assert responses["links"][1]["items"][0]["status_code"] == 404
    assert responses["resources"][1]["items"][0]["size_bytes"] == 123
    assert responses["resources"][1]["items"][0]["title"] == "Missing"
    assert responses["status_codes"][1]["items"][0]["status_code"] == 200
    assert responses["visualization"][1]["nodes"][0]["label"] == "Home"
    assert responses["visualization"][1]["total_nodes"] == 1
    assert responses["visualization"][1]["total_edges"] == 0
    assert responses["visualization"][1]["truncated"] is False


def test_checkpoint_results_can_be_exported_after_pause() -> None:
    service, _, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
        run_id = created["run_id"]
        repository.runs[("test-org", "example", run_id)].status = "paused"
        repository.checkpoints[run_id] = {
            "pages": [
                {
                    "url": "https://example.com/partial",
                    "final_url": "https://example.com/partial",
                    "status_code": 200,
                    "title": "Partial page",
                    "content_type": "text/html",
                }
            ]
        }
        status_code, _, body = asyncio.run(
            api_raw_request(
                "GET",
                f"/api/v1/projects/example/audit-runs/{run_id}/export?dataset=pages&format=json",
            )
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 200
    assert json.loads(body)[0]["title"] == "Partial page"


def test_visualization_does_not_mark_unrenderable_links_as_truncated() -> None:
    service, _, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
    finally:
        app.dependency_overrides.clear()
    run_id = created["run_id"]
    repository.visualization_rows[run_id] = (
        [
            (
                SimpleNamespace(
                    page_id=1,
                    final_url="https://example.com/",
                    title="Home",
                    status_code=200,
                    depth=0,
                ),
                0,
            )
        ],
        [
            (
                SimpleNamespace(
                    id=1,
                    target_url="https://example.com/app.js",
                    placement="body",
                    anchor_text="",
                ),
                "https://example.com/",
            )
        ],
        1,
        1,
    )

    response = asyncio.run(service.visualization("example", run_id))

    assert response.nodes[0].label == "Home"
    assert response.edges == []
    assert response.total_edges == 1
    assert response.truncated is False


def test_archive_current_run_falls_back_to_latest_unarchived_run() -> None:
    service, _, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, first = asyncio.run(api_request("POST", "/api/v1/projects/example/audit-runs", json={}))
        repository.complete(
            first["run_id"],
            {
                "page_count": 1,
                "health_score": 70,
                "errors": 1,
                "warnings": 0,
                "notices": 0,
                "rendered_pages": 0,
            },
        )
        _, second = asyncio.run(api_request("POST", "/api/v1/projects/example/audit-runs", json={}))
        repository.complete(
            second["run_id"],
            {
                "page_count": 2,
                "health_score": 95,
                "errors": 0,
                "warnings": 0,
                "notices": 1,
                "rendered_pages": 0,
            },
        )

        status_code, _ = asyncio.run(
            api_request(
                "POST",
                f"/api/v1/projects/example/audit-runs/{second['run_id']}/archive",
            )
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 200
    assert repository.projects[("test-org", "example")].audit_run_id == first["run_id"]


def test_exports_issues_as_csv_json_and_xml() -> None:
    service, _, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
        run_id = created["run_id"]
        repository.issues[run_id] = [
            AuditIssue(
                id=1,
                run_id=run_id,
                page_id=1,
                url="https://example.com/a",
                severity="error",
                category="content",
                code="missing_title",
                issue="缺少标题",
                details="页面没有 title",
            )
        ]
        results = {
            format_value: asyncio.run(
                api_raw_request(
                    "GET",
                    (
                        f"/api/v1/projects/example/audit-runs/{run_id}/export"
                        f"?dataset=issues&format={format_value}"
                    ),
                )
            )
            for format_value in ("csv", "json", "xml")
        }
    finally:
        app.dependency_overrides.clear()

    assert results["csv"][0] == 200
    assert "text/csv" in results["csv"][1]
    assert "missing_title" in results["csv"][2]
    assert results["json"][0] == 200
    assert "application/json" in results["json"][1]
    assert '"code": "missing_title"' in results["json"][2]
    assert results["xml"][0] == 200
    assert "application/xml" in results["xml"][1]
    assert "<code>missing_title</code>" in results["xml"][2]


def test_exports_pages_and_links_as_csv_json_and_xml() -> None:
    service, _, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
        run_id = created["run_id"]
        results = {
            (dataset, format_value): asyncio.run(
                api_raw_request(
                    "GET",
                    (
                        f"/api/v1/projects/example/audit-runs/{run_id}/export"
                        f"?dataset={dataset}&format={format_value}"
                    ),
                )
            )
            for dataset in ("pages", "links")
            for format_value in ("csv", "json", "xml")
        }
    finally:
        app.dependency_overrides.clear()

    for dataset in ("pages", "links"):
        assert results[(dataset, "csv")][0] == 200
        assert "text/csv" in results[(dataset, "csv")][1]
        assert results[(dataset, "json")][0] == 200
        assert "application/json" in results[(dataset, "json")][1]
        assert results[(dataset, "xml")][0] == 200
        assert "application/xml" in results[(dataset, "xml")][1]

    assert repository.page_queries == [(1, 1000, "", None, None)] * 3
    assert repository.link_queries == [(1, 1000, "", None, None)] * 3


def test_export_links_reads_all_batches_without_truncation() -> None:
    service, _, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
        run_id = created["run_id"]
        repository.link_rows[run_id] = [
            (
                SimpleNamespace(
                    id=index,
                    target_url=f"https://example.com/page-{index}",
                    anchor_text=f"Page {index}",
                    target_status=200,
                    placement="body",
                    is_internal=True,
                    rel="",
                    target_domain="example.com",
                    in_navigation=False,
                ),
                "https://example.com/",
            )
            for index in range(1001)
        ]
        status_code, _, body = asyncio.run(
            api_raw_request(
                "GET",
                (f"/api/v1/projects/example/audit-runs/{run_id}/export?dataset=links&format=json"),
            )
        )
    finally:
        app.dependency_overrides.clear()

    exported = json.loads(body)
    assert status_code == 200
    assert len(exported) == 1001
    assert exported[-1]["target_url"] == "https://example.com/page-1000"
    assert repository.link_queries == [
        (1, 1000, "", None, None),
        (2, 1000, "", None, None),
    ]


def test_export_issues_keeps_groups_across_batch_boundaries() -> None:
    service, _, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
        run_id = created["run_id"]
        repository.issues[run_id] = [
            AuditIssue(
                id=index,
                run_id=run_id,
                page_id=index,
                url=f"https://example.com/page-{index}",
                severity="warning",
                category="content",
                code="duplicate_title",
                issue="Duplicate title",
                details="Pages share the same title",
                similarity=1,
            )
            for index in range(1001)
        ]
        status_code, _, body = asyncio.run(
            api_raw_request(
                "GET",
                (f"/api/v1/projects/example/audit-runs/{run_id}/export?dataset=issues&format=json"),
            )
        )
    finally:
        app.dependency_overrides.clear()

    exported = json.loads(body)
    assert status_code == 200
    assert len(exported) == 1
    assert exported[0]["affected_count"] == 1001
    assert exported[0]["urls"][-1] == "https://example.com/page-999"
    assert repository.issue_group_queries == [(1, 1000)]


def test_issue_grouping_ignores_page_specific_details() -> None:
    service, _, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
        run_id = created["run_id"]
        repository.issues[run_id] = [
            AuditIssue(
                id=index,
                run_id=run_id,
                page_id=index,
                url=f"https://example.com/page-{index}",
                severity="warning",
                category="SEO",
                code="title_too_long",
                issue="Title Too Long",
                details=f"Title is {70 + index} characters",
            )
            for index in (1, 2)
        ]
        status_code, body = asyncio.run(
            api_request(
                "GET",
                f"/api/v1/projects/example/audit-runs/{run_id}/issues",
            )
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 200
    assert body["total"] == 1
    assert body["items"][0]["affected_count"] == 2
    assert len(body["items"][0]["raw"]["detail_examples"]) == 2


def test_spreadsheet_exports_neutralize_formula_cells() -> None:
    assert spreadsheet_cell('=HYPERLINK("https://example.test")') == (
        '\'=HYPERLINK("https://example.test")'
    )
    assert spreadsheet_cell("  +1+1") == "'  +1+1"
    assert spreadsheet_cell("Normal title") == "Normal title"


def test_spreadsheet_rows_split_oversized_issue_url_lists() -> None:
    urls = [f"https://example.com/page-{index:04d}" for index in range(5000)]
    related_urls = [f"https://example.com/related-{index:04d}" for index in range(5000)]
    variants = spreadsheet_row_variants(
        {
            "id": "duplicate_content:warning",
            "affected_count": len(urls),
            "urls": urls,
            "raw": {
                "related_urls": related_urls,
                "max_similarity": 0.9,
                "detail_examples": ["Duplicate page"],
            },
        }
    )

    assert len(variants) > 1
    assert [url for row in variants for url in row["urls"]] == urls
    assert [
        url for row in variants for url in row["raw"]["related_urls"]
    ] == related_urls
    assert sum(row["affected_count"] for row in variants) == len(urls)
    for row in variants:
        assert len(spreadsheet_cell(row["urls"])) <= 32_767
        assert len(spreadsheet_cell(row["raw"])) <= 32_767


def test_exports_real_xlsx_workbook() -> None:
    service, _, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
        run_id = created["run_id"]
        repository.issues[run_id] = [
            AuditIssue(
                id=1,
                run_id=run_id,
                page_id=1,
                url="https://example.com/a",
                severity="error",
                category="content",
                code="missing_title",
                issue="缺少标题",
                details="页面没有 title",
            )
        ]
        result = asyncio.run(
            api_binary_request(
                "GET",
                (f"/api/v1/projects/example/audit-runs/{run_id}/export?dataset=issues&format=xlsx"),
            )
        )
    finally:
        app.dependency_overrides.clear()

    assert result[0] == 200
    assert "spreadsheetml.sheet" in result[1]
    assert f"audit-{run_id}-issues.xlsx" in result[2]
    assert result[3].startswith(b"PK")
    workbook = load_workbook(io.BytesIO(result[3]), read_only=True)
    rows = list(workbook["Audit"].iter_rows(values_only=True))
    assert "missing_title" in rows[1]


def test_pagespeed_returns_persisted_scores_and_metrics() -> None:
    service, _, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
        run_id = created["run_id"]
        repository.pagespeed[run_id] = [
            PageSpeedResult(
                id=1,
                run_id=run_id,
                url="https://example.com",
                strategy="mobile",
                performance_score=88,
                accessibility_score=95,
                best_practices_score=91,
                seo_score=100,
                metrics={"largest_contentful_paint": 2.4},
                error=None,
                analyzed_at=datetime.now(UTC),
            )
        ]
        status_code, body = asyncio.run(
            api_request(
                "GET",
                f"/api/v1/projects/example/audit-runs/{run_id}/pagespeed",
            )
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 200
    assert body["total"] == 1
    assert body["items"][0]["strategy"] == "mobile"
    assert body["items"][0]["performance_score"] == 88
    assert body["items"][0]["metrics"]["largest_contentful_paint"] == 2.4


def test_pagespeed_all_errors_are_failed_in_list_and_detail() -> None:
    service, _, repository = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        _, created = asyncio.run(
            api_request(
                "POST",
                "/api/v1/projects/example/audit-runs",
                json={"enable_pagespeed": True},
            )
        )
        run_id = created["run_id"]
        repository.complete(
            run_id,
            {
                "page_count": 1,
                "health_score": 100,
                "errors": 0,
                "warnings": 0,
                "notices": 0,
                "rendered_pages": 0,
            },
        )
        repository.pagespeed[run_id] = [
            PageSpeedResult(
                id=index,
                run_id=run_id,
                url="https://example.com",
                strategy=strategy,
                metrics={},
                error="quota exceeded",
                analyzed_at=datetime.now(UTC),
            )
            for index, strategy in enumerate(("mobile", "desktop"), start=1)
        ]
        list_status, list_body = asyncio.run(
            api_request("GET", "/api/v1/projects/example/audit-runs")
        )
        detail_status, detail_body = asyncio.run(
            api_request(
                "GET",
                f"/api/v1/projects/example/audit-runs/{run_id}",
            )
        )
    finally:
        app.dependency_overrides.clear()

    assert list_status == detail_status == 200
    assert list_body["items"][0]["pagespeed"]["status"] == "failed"
    assert detail_body["pagespeed"]["status"] == "failed"
    assert detail_body["pagespeed"]["message"]


def test_launch_failure_keeps_database_run_queued_for_retry() -> None:
    service, _, repository = build_service(launch_error=RuntimeError("offline"))
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        status_code, _ = asyncio.run(
            api_request("POST", "/api/v1/projects/example/audit-runs", json={})
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 202
    run = next(iter(repository.runs.values()))
    assert run.status == "queued"
    assert run.can_resume is False


def test_directory_scope_requires_directory() -> None:
    service, _, _ = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        status_code, _ = asyncio.run(
            api_request(
                "POST",
                "/api/v1/projects/example/audit-runs",
                json={"scope": "directory"},
            )
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 422


def test_audit_page_limit_cannot_exceed_5000() -> None:
    service, _, _ = build_service()
    app.dependency_overrides[get_audit_service] = lambda: service
    try:
        status_code, _ = asyncio.run(
            api_request(
                "POST",
                "/api/v1/projects/example/audit-runs",
                json={"max_pages": 5001},
            )
        )
    finally:
        app.dependency_overrides.clear()

    assert status_code == 422
